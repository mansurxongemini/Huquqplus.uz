import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlmodel import Session, select
from src.database.mysql import engine
from src.models.inquiry import Inquiry, InquiryStatus
from src.models.appointment import Appointment
from src.models.lawyer import OfficialLawyer
from src.models.volunteer import Volunteer
from src.models.user import User


def generate_staff_kpi_report(date_range: tuple[datetime, datetime] | None = None) -> str:
    """
    Generates a 4-sheet analytical Excel report:
    1. Specialists KPI Ranking (Volunteers & Lawyers)
    2. Registered Volunteers Leaderboard & Profiles
    3. Detailed Inquiries & User Ratings
    4. Lawyer Appointments & Ratings
    """
    with Session(engine) as session:
        # 1. Fetch Inquiries
        inquiry_stmt = select(Inquiry)
        if date_range:
            start_date, end_date = date_range
            inquiry_stmt = inquiry_stmt.where(Inquiry.created_at >= start_date, Inquiry.created_at <= end_date)
        inquiries = session.exec(inquiry_stmt).all()

        # 2. Fetch Appointments
        appointment_stmt = select(Appointment)
        if date_range:
            start_date, end_date = date_range
            appointment_stmt = appointment_stmt.where(Appointment.created_at >= start_date, Appointment.created_at <= end_date)
        appointments = session.exec(appointment_stmt).all()

        # 3. Fetch Official Lawyers & Volunteers & Users
        official_lawyers = session.exec(select(OfficialLawyer)).all()
        lawyer_map = {l.user_id: l for l in official_lawyers}

        volunteers = session.exec(select(Volunteer)).all()
        vol_map = {v.user_id: v for v in volunteers}

        users = session.exec(select(User)).all()
        user_map = {u.user_id: u for u in users}

    wb = Workbook()

    # Styling definitions
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=14, bold=True, color="1F4E78")
    sub_title_font = Font(name="Calibri", size=10, italic=True, color="595959")
    bold_font = Font(name="Calibri", size=11, bold=True)
    regular_font = Font(name="Calibri", size=11)
    total_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # =========================================================================
    # SHEET 1: Mutaxassislar KPI Reytingi (Umumiy)
    # =========================================================================
    ws1 = wb.active
    ws1.title = "Mutaxassislar KPI"
    ws1.views.sheetView[0].showGridLines = True

    ws1.cell(row=1, column=1, value="HuquqPlus — Mutaxassislar (Yurist va Volontyorlar) KPI Tahlili").font = title_font
    report_period = "Barcha davr uchun" if not date_range else f"Davr: {date_range[0].strftime('%Y-%m-%d')} dan {date_range[1].strftime('%Y-%m-%d')} gacha"
    ws1.cell(row=2, column=1, value=f"{report_period} | Yaratilgan vaqt: {datetime.now().strftime('%Y-%m-%d %H:%M')}").font = sub_title_font

    headers1 = [
        "№", "Mutaxassis F.I.Sh.", "Telegram ID", "Telefon", "Rol / Maqomi", "O'qish / Ish joyi",
        "Jami KPI Ball", "Javoblar soni", "O'rtacha baho (⭐)", "O'rtacha javob tezligi",
        "5⭐ baholar", "4⭐ baholar", "1-2⭐ baholar", "Qabullar soni"
    ]

    for col_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=4, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    specialist_stats = {}

    # Initialize from official lawyers
    for l in official_lawyers:
        specialist_stats[l.user_id] = {
            "name": l.full_name,
            "phone": l.phone or "—",
            "role": "Yurist",
            "study_work": l.specialization or "Huquqshunos",
            "inquiries_count": 0,
            "inquiries_ratings": [],
            "five_star": 0,
            "four_star": 0,
            "negative_ratings": 0,
            "response_times": [],
            "fast_responses": 0,
            "appointments_count": 0,
            "appointments_ratings": [],
        }

    # Initialize from registered volunteers
    for v in volunteers:
        if v.user_id not in specialist_stats:
            specialist_stats[v.user_id] = {
                "name": v.full_name,
                "phone": v.phone_number or "—",
                "role": "Volontyor",
                "study_work": v.study_or_work or "—",
                "inquiries_count": 0,
                "inquiries_ratings": [],
                "five_star": 0,
                "four_star": 0,
                "negative_ratings": 0,
                "response_times": [],
                "fast_responses": 0,
                "appointments_count": 0,
                "appointments_ratings": [],
            }

    # Process inquiries
    for inq in inquiries:
        resp_id = inq.responder_id
        if not resp_id:
            continue

        if resp_id not in specialist_stats:
            user_obj = user_map.get(resp_id)
            name = f"{user_obj.first_name or ''} {user_obj.last_name or ''}".strip() if user_obj else f"Mutaxassis {resp_id}"
            specialist_stats[resp_id] = {
                "name": name or f"ID: {resp_id}",
                "phone": user_obj.phone if user_obj else "—",
                "role": "Volontyor / Foydalanuvchi",
                "study_work": "—",
                "inquiries_count": 0,
                "inquiries_ratings": [],
                "five_star": 0,
                "four_star": 0,
                "negative_ratings": 0,
                "response_times": [],
                "fast_responses": 0,
                "appointments_count": 0,
                "appointments_ratings": [],
            }

        st = specialist_stats[resp_id]
        st["inquiries_count"] += 1

        if inq.rating is not None:
            st["inquiries_ratings"].append(inq.rating)
            if inq.rating == 5:
                st["five_star"] += 1
            elif inq.rating == 4:
                st["four_star"] += 1
            elif inq.rating <= 2:
                st["negative_ratings"] += 1

        if inq.created_at and inq.replied_at:
            try:
                rep_time = inq.replied_at if isinstance(inq.replied_at, datetime) else datetime.fromisoformat(str(inq.replied_at))
                created_time = inq.created_at if isinstance(inq.created_at, datetime) else datetime.fromisoformat(str(inq.created_at))
                diff_min = (rep_time - created_time).total_seconds() / 60.0
                if diff_min >= 0:
                    st["response_times"].append(diff_min)
                    if diff_min <= 60.0:
                        st["fast_responses"] += 1
            except Exception:
                pass

    # Process appointments
    for app in appointments:
        lawyer_id = app.lawyer_id
        if not lawyer_id or lawyer_id not in specialist_stats:
            continue
        st = specialist_stats[lawyer_id]
        if app.status == "completed":
            st["appointments_count"] += 1
            if app.rating:
                st["appointments_ratings"].append(app.rating)

    # Calculate KPI & format
    processed_specs = []
    for uid, data in specialist_stats.items():
        kpi = (
            (data["inquiries_count"] * 10)
            + (data["five_star"] * 15)
            + (data["four_star"] * 10)
            - (data["negative_ratings"] * 10)
            + (data["fast_responses"] * 5)
            + (data["appointments_count"] * 20)
        )
        avg_resp = (sum(data["response_times"]) / len(data["response_times"])) if data["response_times"] else 0.0
        if avg_resp == 0:
            resp_str = "—"
        elif avg_resp < 60:
            resp_str = f"{int(avg_resp)} daqiqa"
        else:
            resp_str = f"{int(avg_resp // 60)}s {int(avg_resp % 60)}d"

        data["uid"] = uid
        data["kpi_score"] = max(kpi, 0)
        data["avg_response_str"] = resp_str
        processed_specs.append(data)

    processed_specs.sort(key=lambda x: (x["kpi_score"], x["inquiries_count"]), reverse=True)

    current_row = 5
    for idx, data in enumerate(processed_specs, 1):
        inq_avg = round(sum(data["inquiries_ratings"]) / len(data["inquiries_ratings"]), 2) if data["inquiries_ratings"] else "—"

        row_vals = [
            idx,
            data["name"],
            data["uid"],
            data["phone"],
            data["role"],
            data["study_work"],
            data["kpi_score"],
            data["inquiries_count"],
            inq_avg,
            data["avg_response_str"],
            data["five_star"],
            data["four_star"],
            data["negative_ratings"],
            data["appointments_count"]
        ]

        for col_idx, val in enumerate(row_vals, 1):
            c = ws1.cell(row=current_row, column=col_idx, value=val)
            c.font = regular_font
            c.border = thin_border
            c.alignment = align_center if col_idx in [1, 3, 4, 5, 9, 10] else (align_right if col_idx in [7, 8, 11, 12, 13, 14] else align_left)

        current_row += 1

    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # =========================================================================
    # SHEET 2: Volontyorlar Reytingi (Faqat ro'yxatdan o'tgan volontyorlar)
    # =========================================================================
    ws2 = wb.create_sheet(title="Volontyorlar Reytingi")
    ws2.views.sheetView[0].showGridLines = True

    ws2.cell(row=1, column=1, value="HuquqPlus — Rasmiy Volontyorlar Leaderboard va Faoliyat Tahlili").font = title_font
    ws2.cell(row=2, column=1, value=f"Ro'yxatdan o'tgan barcha volontyorlar reytingi | {datetime.now().strftime('%Y-%m-%d %H:%M')}").font = sub_title_font

    headers2 = [
        "O'rin", "Volontyor F.I.Sh.", "Telegram ID", "Telefon raqami", "O'qish / Ish joyi", "Tug'ilgan yili",
        "Jami KPI Ball", "Javoblar soni", "O'rtacha baho (⭐)", "O'rtacha javob tezligi",
        "5⭐ baholar", "4⭐ baholar", "1-2⭐ baholar", "Ro'yxatdan o'tgan vaqti"
    ]

    for col_idx, h in enumerate(headers2, 1):
        cell = ws2.cell(row=4, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    # Filter only registered volunteers
    vol_specs = [s for s in processed_specs if s["uid"] in vol_map]
    vol_specs.sort(key=lambda x: (x["kpi_score"], x["inquiries_count"]), reverse=True)

    r_row = 5
    for idx, data in enumerate(vol_specs, 1):
        v_obj = vol_map.get(data["uid"])
        inq_avg = round(sum(data["inquiries_ratings"]) / len(data["inquiries_ratings"]), 2) if data["inquiries_ratings"] else "—"
        reg_time = v_obj.created_at.strftime('%Y-%m-%d %H:%M') if v_obj and v_obj.created_at else "—"
        b_year = v_obj.birth_year if v_obj else "—"

        row_v = [
            idx,
            data["name"],
            data["uid"],
            data["phone"],
            data["study_work"],
            b_year,
            data["kpi_score"],
            data["inquiries_count"],
            inq_avg,
            data["avg_response_str"],
            data["five_star"],
            data["four_star"],
            data["negative_ratings"],
            reg_time
        ]

        for col_idx, val in enumerate(row_v, 1):
            c = ws2.cell(row=r_row, column=col_idx, value=val)
            c.font = regular_font
            c.border = thin_border
            c.alignment = align_center if col_idx in [1, 3, 4, 6, 9, 10, 14] else (align_right if col_idx in [7, 8, 11, 12, 13] else align_left)

        r_row += 1

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # =========================================================================
    # SHEET 3: Murojaatlar tafsiloti
    # =========================================================================
    ws3 = wb.create_sheet(title="Murojaatlar tafsiloti")
    ws3.views.sheetView[0].showGridLines = True

    headers3 = [
        "ID", "Sana-vaqt", "Bo'lim", "Fuqaro F.I.Sh.", "Fuqaro Tel",
        "Savol matni", "Javob bergan mutaxassis", "Mutaxassis ID", "Rol",
        "Javob matni", "Javob berilgan vaqt", "Fuqaro bahosi (1-5⭐)", "Fuqaro izohi"
    ]

    for col_idx, h in enumerate(headers3, 1):
        cell = ws3.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    for r_idx, inq in enumerate(inquiries, 2):
        u_obj = user_map.get(inq.user_id) if inq.user_id else None
        citizen_name = f"{u_obj.first_name or ''} {u_obj.last_name or ''}".strip() if u_obj else f"User {inq.user_id}"
        citizen_phone = u_obj.phone if u_obj else "—"

        resp_stat = specialist_stats.get(inq.responder_id)
        resp_name = resp_stat["name"] if resp_stat else (f"ID: {inq.responder_id}" if inq.responder_id else "Javob berilmagan")
        resp_role = resp_stat["role"] if resp_stat else "—"

        created_str = inq.created_at.strftime('%Y-%m-%d %H:%M') if inq.created_at else "—"
        replied_str = str(inq.replied_at)[:16] if inq.replied_at else "—"
        rating_str = f"⭐ {inq.rating}" if inq.rating else "Baholanmagan"

        row3 = [
            inq.id,
            created_str,
            inq.section_name or "Boshqa",
            citizen_name,
            citizen_phone or "—",
            inq.question[:300] if inq.question else "—",
            resp_name,
            inq.responder_id or "—",
            resp_role,
            (inq.answer[:300] if inq.answer else "—"),
            replied_str,
            rating_str,
            inq.feedback or "—"
        ]

        for col_idx, val in enumerate(row3, 1):
            c = ws3.cell(row=r_idx, column=col_idx, value=val)
            c.font = regular_font
            c.border = thin_border
            c.alignment = align_center if col_idx in [1, 2, 5, 8, 9, 11, 12] else align_left

    for col in ws3.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws3.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    # =========================================================================
    # SHEET 4: Yurist Qabullari
    # =========================================================================
    ws4 = wb.create_sheet(title="Yurist Qabullari")
    ws4.views.sheetView[0].showGridLines = True

    headers4 = [
        "Ariza ID", "Ariza sanasi", "Fuqaro F.I.Sh.", "Telefon",
        "Qabul turi", "Muammo mavzusi", "Biriktirilgan yurist", "Yurist ID",
        "Holati", "Suhbat yakunlangan vaqt", "Fuqaro bahosi (1-5⭐)", "Fuqaro izohi"
    ]

    for col_idx, h in enumerate(headers4, 1):
        cell = ws4.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center

    for r_idx, app in enumerate(appointments, 2):
        u_obj = user_map.get(app.user_id) if app.user_id else None
        citizen_name = f"{u_obj.first_name or ''} {u_obj.last_name or ''}".strip() if u_obj else f"User {app.user_id}"
        citizen_phone = u_obj.phone if u_obj else "—"

        lawyer_obj = lawyer_map.get(app.lawyer_id)
        lawyer_name = lawyer_obj.full_name if lawyer_obj else (f"ID: {app.lawyer_id}" if app.lawyer_id else "Biriktirilmagan")

        status_uz = {
            "pending": "Kutilmoqda",
            "accepted": "Qabul qilingan",
            "completed": "Yakunlangan",
            "cancelled": "Bekor qilingan"
        }.get(app.status, app.status)

        type_uz = "Masofaviy (Online)" if app.appointment_type == "remote_appointment" else "Yuzma-yuz (Ofis)"
        created_str = app.created_at.strftime('%Y-%m-%d %H:%M') if app.created_at else "—"
        completed_str = app.completed_at.strftime('%Y-%m-%d %H:%M') if app.completed_at else "—"
        rating_str = f"⭐ {app.rating}" if app.rating else "Baholanmagan"

        row4 = [
            app.id,
            created_str,
            citizen_name,
            citizen_phone or "—",
            type_uz,
            app.problem_description[:300] if app.problem_description else "—",
            lawyer_name,
            app.lawyer_id or "—",
            status_uz,
            completed_str,
            rating_str,
            app.feedback or "—"
        ]

        for col_idx, val in enumerate(row4, 1):
            c = ws4.cell(row=r_idx, column=col_idx, value=val)
            c.font = regular_font
            c.border = thin_border
            c.alignment = align_center if col_idx in [1, 2, 4, 5, 8, 9, 10, 11] else align_left

    for col in ws4.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = col[0].column_letter
        ws4.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

    # Save to file
    today_str = datetime.now().strftime('%Y-%m-%d_%H-%M')
    target_dir = "/app/user_data/reports"
    os.makedirs(target_dir, exist_ok=True)
    file_path = f"{target_dir}/staff_kpi_report_{today_str}.xlsx"
    wb.save(file_path)

    return file_path
